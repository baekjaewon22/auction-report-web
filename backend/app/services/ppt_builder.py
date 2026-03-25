# -*- coding: utf-8 -*-
"""
PPT 생성 서비스
- PPTM 템플릿 로드
- 데이터 채우기 (1번/3번 슬라이드)
- 사진 삽입
- 슬라이드 복제/이동
- 노란 박스 찾기 → 이미지 교체
- VBA 보존 저장
"""

import os
import re
import logging
import tempfile
import zipfile
from io import BytesIO
from copy import deepcopy
from typing import Optional, Tuple, Dict

import requests
from pptx import Presentation

from ..core.config import settings, CAPTURE_DIR
from ..core.utils import track_file
from .capturer import trim_white_margin

logger = logging.getLogger(__name__)

TARGET_SLIDE_INDEX = 2  # 0-base → 3번 슬라이드


# ============================================================
# AltText / 토큰
# ============================================================
def get_alt_text(shape) -> str:
    try:
        cNvPr = shape._element.xpath(".//p:cNvPr")[0]
        return cNvPr.get("descr") or cNvPr.get("title") or ""
    except Exception:
        return ""


_NUM_RE = re.compile(r"(\d[\d,]*)")


def parse_token(alt: str) -> Tuple[str, Dict[str, str]]:
    alt = (alt or "").strip()
    if not alt.startswith("EXCEL_"):
        return "", {}
    parts = alt.split("|")
    head = parts[0]
    opts: Dict[str, str] = {}
    for p in parts[1:]:
        if "=" in p:
            k, v = p.split("=", 1)
            opts[k.strip().upper()] = v
    if ":" not in head:
        return "", {}
    kind, payload = head.split(":", 1)
    opts["PAYLOAD"] = payload.strip()
    return kind.strip().upper(), opts


# ============================================================
# 텍스트 유틸
# ============================================================
def replace_first_number_preserve_runs(text_frame, new_number: str) -> bool:
    runs, full = [], ""
    for p in text_frame.paragraphs:
        for r in p.runs:
            runs.append(r)
            full += r.text

    m = _NUM_RE.search(full)
    if not m:
        return False

    start, end = m.span(1)
    idx, replaced = 0, False
    for r in runs:
        t = r.text or ""
        rs, re_ = idx, idx + len(t)
        if re_ <= start or rs >= end:
            idx = re_
            continue
        a = max(start - rs, 0)
        b = min(end - rs, len(t))
        if not replaced:
            r.text = t[:a] + new_number + t[b:]
            replaced = True
        else:
            r.text = t[:a] + t[b:]
        idx = re_
    return True


def set_text_keep_style(text_frame, new_text: str):
    p = text_frame.paragraphs[0] if text_frame.paragraphs else text_frame.add_paragraph()
    font_tpl = p.runs[0].font if p.runs else None
    p.clear()
    r = p.add_run()
    r.text = new_text
    if font_tpl:
        r.font.name = font_tpl.name
        r.font.size = font_tpl.size
        r.font.bold = font_tpl.bold
        r.font.italic = font_tpl.italic
        r.font.underline = font_tpl.underline
        try:
            r.font.color.rgb = font_tpl.color.rgb
        except Exception:
            pass


_WON_RE = re.compile(r"(\d[\d,]*)\s*원")


def remove_won_unit_in_slide(slide):
    for shp in slide.shapes:
        if not getattr(shp, "has_text_frame", False):
            continue
        tf = shp.text_frame
        old = "\n".join([p.text for p in tf.paragraphs])
        new = _WON_RE.sub(r"\1", old)
        if new != old:
            tf.clear()
            for i, line in enumerate(new.split("\n")):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line


# ============================================================
# 슬라이드 검색
# ============================================================
def find_slide_by_keyword(prs: Presentation, keyword: str):
    hangul_re = re.compile(r"[가-힣]")
    for idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            text = (shape.text or "").strip()
            if not text or keyword not in text:
                continue
            rest = text.replace(keyword, "")
            if hangul_re.search(rest):
                continue
            return slide
    return None


def find_slide_by_note_key(prs: Presentation, key: str):
    for s in prs.slides:
        try:
            notes = s.notes_slide.notes_text_frame.text or ""
            if key in notes:
                return s
        except Exception:
            continue
    return None


def find_slide_index_by_note_key(prs: Presentation, key: str) -> int:
    for idx, s in enumerate(prs.slides):
        try:
            notes = s.notes_slide.notes_text_frame.text or ""
            if key in notes:
                return idx
        except Exception:
            continue
    return -1


# ============================================================
# 노란 박스 찾기
# ============================================================
def find_yellow_box(slide):
    candidates = []
    for shape in slide.shapes:
        text = ""
        if hasattr(shape, "text"):
            try:
                text = shape.text or ""
            except Exception:
                text = ""
        if text.strip():
            continue
        try:
            fill = shape.fill
        except Exception:
            continue
        if not fill or fill.type is None:
            continue
        try:
            area = shape.width * shape.height
        except Exception:
            continue
        candidates.append((area, shape))
    if candidates:
        candidates.sort(key=lambda x: x[0], reverse=True)
        return candidates[0][1]
    return None


def find_yellow_boxes_left_to_right(slide, limit=2):
    candidates = []
    for shape in slide.shapes:
        try:
            if hasattr(shape, "text") and (shape.text or "").strip():
                continue
        except Exception:
            pass
        try:
            fill = shape.fill
            if not fill or fill.type is None:
                continue
        except Exception:
            continue
        try:
            area = shape.width * shape.height
        except Exception:
            continue
        candidates.append((shape.left, -area, shape))
    candidates.sort(key=lambda x: (x[0], x[1]))
    return [c[2] for c in candidates[:limit]]


# ============================================================
# 슬라이드 복제/이동
# ============================================================
def duplicate_slide(prs: Presentation, slide):
    new_slide = prs.slides.add_slide(slide.slide_layout)
    for shape in slide.shapes:
        new_el = deepcopy(shape._element)
        new_slide.shapes._spTree.insert_element_before(new_el, "p:extLst")
    return new_slide


def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slide_id = xml_slides[old_index]
    xml_slides.remove(slide_id)
    if new_index > old_index:
        new_index -= 1
    xml_slides.insert(new_index, slide_id)


# ============================================================
# 데이터 채우기
# ============================================================
def fill_slide_with_data(prs: Presentation, data: dict):
    for idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            placeholder = (shape.text or "").strip()
            if not placeholder:
                continue
            key = None
            if placeholder.startswith("{") and placeholder.endswith("}"):
                key = placeholder[1:-1].strip()
            elif placeholder.startswith("[") and placeholder.endswith("]"):
                key = placeholder[1:-1].strip()
            if not key or key not in data:
                continue
            if key == "address":
                main_addr = data.get("address", "")
                old_addr = data.get("address_old", "")
                if idx == 0 and old_addr:
                    shape.text = f"{main_addr}\n{old_addr}"
                else:
                    shape.text = main_addr
            else:
                shape.text = data[key]


def insert_main_photo(prs: Presentation, photo_url: str):
    if not photo_url:
        return
    try:
        resp = requests.get(photo_url, timeout=15)
        resp.raise_for_status()
    except Exception as e:
        logger.warning(f"사진 다운로드 실패: {e}")
        return
    img_bytes = BytesIO(resp.content)
    slide = prs.slides[TARGET_SLIDE_INDEX]
    try:
        picture_ph = slide.placeholders[13]
        picture_ph.insert_picture(img_bytes)
        logger.info("3번 슬라이드에 대표사진 삽입 완료")
    except Exception as e:
        logger.warning(f"사진 placeholder 삽입 실패: {e}")


def replace_placeholders_in_slide(slide, mapping: dict):
    for shp in slide.shapes:
        if not getattr(shp, "has_text_frame", False):
            continue
        tf = shp.text_frame
        old = "\n".join([p.text for p in tf.paragraphs])
        new = old
        for k, v in mapping.items():
            new = new.replace(k, "" if v is None else str(v))
        if new != old:
            tf.clear()
            for i, line in enumerate(new.split("\n")):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line


# ============================================================
# 이미지 삽입 (복제 + 노란박스 교체)
# ============================================================
def insert_images_into_ppt(prs, total_pages, keyword, img_pattern,
                           clone_from_next=False, force_title=False,
                           forced_num="(5)", base_as_last=False):
    if total_pages <= 0:
        return

    base_slide = find_slide_by_keyword(prs, keyword)
    if base_slide is None:
        logger.error(f"'{keyword}' 기준 슬라이드를 찾지 못했습니다.")
        return

    # 제목 텍스트 보존
    base_title_text = ""
    base_forced_num = forced_num
    for shp in base_slide.shapes:
        if not hasattr(shp, "text"):
            continue
        txt = (shp.text or "").strip()
        if txt and keyword in txt:
            base_title_text = re.sub(r"-\d+\s*$", "", txt).strip()
            m = re.search(r"\(\d+\)", base_title_text)
            if m:
                base_forced_num = m.group(0)
            break
    if not base_title_text:
        base_title_text = f"{base_forced_num} {keyword}"

    def _apply_title(dst):
        for shp in dst.shapes:
            if hasattr(shp, "text") and keyword in (shp.text or ""):
                shp.text = base_title_text
                return

    # 템플릿 슬라이드 결정
    template_slide = base_slide
    base_index = prs.slides.index(base_slide)

    if clone_from_next:
        if keyword in ["건축물대장", "토지이용계획", "전자지도", "위성지도", "위치도"]:
            tpl_key = "SLIDE_KEY=TEMPLATE_OBJ_STATUS_PAGE"
        elif keyword in ["매각물건명세서", "등기사항 요약", "현황조사서", "임차인 현황", "등기부 현황"] or "등기부 현황" in keyword:
            tpl_key = "SLIDE_KEY=TEMPLATE_RIGHT_ANALYSIS_PAGE"
        else:
            tpl_key = ""

        if tpl_key:
            tpl_idx = find_slide_index_by_note_key(prs, tpl_key)
            if tpl_idx >= 0:
                template_slide = prs.slides[tpl_idx]

    # 슬라이드-페이지 매핑
    slide_for_page = {}
    if not base_as_last:
        slide_for_page[1] = base_slide
        for page in range(2, total_pages + 1):
            new_s = duplicate_slide(prs, template_slide)
            slide_for_page[page] = new_s
            _apply_title(new_s)
        for page in range(2, total_pages + 1):
            old_idx = prs.slides.index(slide_for_page[page])
            move_slide(prs, old_idx, base_index + (page - 1))
    else:
        for page in range(1, total_pages):
            new_s = duplicate_slide(prs, template_slide)
            slide_for_page[page] = new_s
            _apply_title(new_s)
        slide_for_page[total_pages] = base_slide
        for page in range(1, total_pages):
            old_idx = prs.slides.index(slide_for_page[page])
            move_slide(prs, old_idx, base_index + (page - 1))

    # 이미지 삽입
    for page in range(1, total_pages + 1):
        slide = slide_for_page.get(page)
        if slide is None:
            continue
        img_path = img_pattern.format(page=page)
        if not os.path.exists(img_path):
            continue
        yellow = find_yellow_box(slide)
        if yellow is None:
            continue
        left, top, width, height = yellow.left, yellow.top, yellow.width, yellow.height
        slide.shapes._spTree.remove(yellow._element)

        trimmed = img_path.replace(".png", "_trim.png")
        try:
            trim_white_margin(img_path, trimmed)
            track_file(trimmed)
            use_path = trimmed
        except Exception:
            use_path = img_path

        slide.shapes.add_picture(use_path, left, top, width=width, height=height)

        # 제목 번호
        if total_pages > 1:
            for shp in slide.shapes:
                if hasattr(shp, "text") and keyword in (shp.text or ""):
                    base = re.sub(r"-\d+\s*$", "", shp.text).strip()
                    shp.text = f"{base}-{page}"
                    break


def insert_single_image(prs, keyword_or_key, image_path, use_note_key=False):
    if not image_path or not os.path.exists(image_path):
        return
    slide = (find_slide_by_note_key(prs, keyword_or_key) if use_note_key
             else find_slide_by_keyword(prs, keyword_or_key))
    if slide is None:
        return
    yellow = find_yellow_box(slide)
    if yellow is None:
        return
    l, t, w, h = yellow.left, yellow.top, yellow.width, yellow.height
    slide.shapes._spTree.remove(yellow._element)

    trimmed = image_path.replace(".png", "_trim.png")
    try:
        trim_white_margin(image_path, trimmed)
        track_file(trimmed)
        use_path = trimmed
    except Exception:
        use_path = image_path
    slide.shapes.add_picture(use_path, l, t, width=w, height=h)


def insert_two_images_location(prs, keyword, left_img, right_img=""):
    slide = find_slide_by_keyword(prs, keyword)
    if slide is None:
        return
    boxes = find_yellow_boxes_left_to_right(slide, limit=2)
    if not boxes:
        return

    def _put(box, img_path):
        if not img_path or not os.path.exists(img_path):
            return
        l, t, w, h = box.left, box.top, box.width, box.height
        slide.shapes._spTree.remove(box._element)
        trimmed = img_path.replace(".png", "_trim.png")
        try:
            trim_white_margin(img_path, trimmed)
            track_file(trimmed)
            use_path = trimmed
        except Exception:
            use_path = img_path
        slide.shapes.add_picture(use_path, l, t, width=w, height=h)

    _put(boxes[0], left_img)
    if len(boxes) > 1:
        _put(boxes[1], right_img)


# ============================================================
# VBA 보존 저장
# ============================================================
def save_pptm_preserve_vba(template_pptm: str, prs: Presentation, out_pptm: str):
    tmp_pptx = tempfile.mktemp(suffix=".pptx")
    prs.save(tmp_pptx)

    with zipfile.ZipFile(template_pptm, "r") as zt:
        vba_bin = zt.read("ppt/vbaProject.bin")

    with zipfile.ZipFile(tmp_pptx, "r") as zi, \
         zipfile.ZipFile(out_pptm, "w", compression=zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            if item.filename == "ppt/vbaProject.bin":
                continue
            zo.writestr(item, zi.read(item.filename))
        zo.writestr("ppt/vbaProject.bin", vba_bin)

    try:
        os.remove(tmp_pptx)
    except Exception:
        pass
